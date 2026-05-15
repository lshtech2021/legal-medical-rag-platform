from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def render_docx(title: str, body: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for para in body.split("\n"):
        if para.strip():
            doc.add_paragraph(para)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def render_pdf(title: str, body: str) -> bytes:
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=letter)
    width, height = letter
    y = height - 50
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, y, title)
    y -= 30
    pdf.setFont("Helvetica", 11)
    for line in body.split("\n"):
        if y < 50:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 11)
        pdf.drawString(50, y, line[:120])
        y -= 16
    pdf.save()
    return output.getvalue()
